import csv
import io
import zipfile
from odoo.exceptions import UserError


def parse_file(payload, filename, headers):
    if filename.lower().endswith(".csv"):
        try:
            values = list(csv.reader(io.StringIO(payload.decode("utf-8-sig"))))
        except (UnicodeError, csv.Error) as e:
            raise UserError("Invalid UTF-8 CSV.") from e
    elif filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as z:
                if sum(f.file_size for f in z.infolist()) > 30 * 1024 * 1024:
                    raise UserError("Expanded XLSX exceeds 30 MB.")
            book = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
            if len(book.sheetnames) != 1:
                raise UserError("Use one worksheet per import.")
            values = []
            for row in book.active.iter_rows():
                if any(c.data_type == "f" for c in row):
                    raise UserError("Formula cells are not accepted.")
                values.append([c.value for c in row])
                if len(values) > 10001:
                    raise UserError("Maximum 10,000 rows.")
            book.close()
        except (ValueError, zipfile.BadZipFile) as e:
            raise UserError("Invalid XLSX file.") from e
    else:
        raise UserError("Upload a CSV or XLSX file.")
    if not values or [str(v or "").strip() for v in values[0]] != headers:
        raise UserError("Headers must match the downloadable template exactly.")
    if len(values) > 10001:
        raise UserError("Maximum 10,000 rows.")
    rows = []
    for number, row in enumerate(values[1:], 2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        if len(row) != len(headers):
            raise UserError(f"Row {number}: incorrect column count.")
        rows.append(
            (number, {h: str(v).strip() if v is not None else "" for h, v in zip(headers, row)})
        )
    if not rows:
        raise UserError("The import has no data rows.")
    return rows
